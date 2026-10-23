"""Unit tests for document converters and Google Drive URL extraction."""

import io
import pytest
from converters.txt_converter import convert_txt_to_markdown
from converters.docx_converter import convert_docx_to_markdown
from converters.gdrive_fetcher import (
    extract_gdrive_id,
    GDriveInvalidUrlError,
    _detect_format,
)


def test_txt_converter():
    """Verify plain text and markdown conversion."""
    raw_text = "John Doe\nSoftware Engineer\nPython, ABAP, SQL"
    res = convert_txt_to_markdown(raw_text)
    assert "John Doe" in res
    assert "ABAP" in res


def test_txt_converter_bytes():
    """Verify bytes decoding with UTF-8."""
    raw_bytes = "Jane Doe\nDeveloper".encode("utf-8")
    res = convert_txt_to_markdown(raw_bytes)
    assert "Jane Doe" in res


def test_docx_converter():
    """Verify Word document parser extracts paragraphs and tables."""
    import docx
    doc = docx.Document()
    doc.add_heading("Alice Candidate", level=1)
    doc.add_paragraph("Full Stack Developer with 2 years of experience.")
    
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Experience"
    table.rows[1].cells[0].text = "Python"
    table.rows[1].cells[1].text = "3 years"

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    md = convert_docx_to_markdown(bio.read())
    assert "# Alice Candidate" in md
    assert "Full Stack Developer" in md
    assert "| Skill | Experience |" in md


def test_gdrive_id_extraction_standard():
    """Verify standard sharing link extraction."""
    url = "https://drive.google.com/file/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/view?usp=sharing"
    file_id, doc_type = extract_gdrive_id(url)
    assert file_id == "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"
    assert doc_type == "file"


def test_gdrive_id_extraction_docs():
    """Verify Google Docs link extraction."""
    url = "https://docs.google.com/document/d/12345abcde67890/edit"
    file_id, doc_type = extract_gdrive_id(url)
    assert file_id == "12345abcde67890"
    assert doc_type == "document"


def test_gdrive_invalid_url():
    """Verify invalid URLs raise GDriveInvalidUrlError."""
    with pytest.raises(GDriveInvalidUrlError):
        extract_gdrive_id("https://example.com/not-gdrive")


def test_format_detector():
    """Verify format detection handles magic bytes and mime types."""
    pdf_bytes = b"%PDF-1.4\n..."
    assert _detect_format(pdf_bytes, "application/pdf", "resume.pdf") == "pdf"
    
    txt_bytes = b"Hello world text"
    assert _detect_format(txt_bytes, "text/plain", "resume.txt") == "txt"
