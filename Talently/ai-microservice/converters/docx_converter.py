"""DOCX (Word Document) to Markdown converter using python-docx."""

import io
import logging
from typing import Union
import docx

logger = logging.getLogger("sap_evaluator.docx_converter")


def convert_docx_to_markdown(docx_source: Union[str, bytes]) -> str:
    """Convert a Word DOCX document (path or raw bytes) to Markdown format.

    Extracts paragraphs, headings, bullet lists, and tables.

    Args:
        docx_source: File path as string or raw bytes of the DOCX document.

    Returns:
        Clean markdown string extracted from the document.

    Raises:
        ValueError: If the file cannot be parsed as a valid DOCX document.
    """
    try:
        if isinstance(docx_source, bytes):
            file_stream = io.BytesIO(docx_source)
            doc = docx.Document(file_stream)
        elif isinstance(docx_source, str):
            doc = docx.Document(docx_source)
        else:
            raise ValueError(f"Invalid docx_source type: {type(docx_source)}")

        md_lines = []

        # Process paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name.lower() if para.style else ""

            if "heading 1" in style_name:
                md_lines.append(f"# {text}")
            elif "heading 2" in style_name:
                md_lines.append(f"## {text}")
            elif "heading 3" in style_name:
                md_lines.append(f"### {text}")
            elif "list" in style_name or "bullet" in style_name:
                md_lines.append(f"- {text}")
            else:
                md_lines.append(text)

        # Process tables if present
        for table in doc.tables:
            if not table.rows:
                continue
            
            md_lines.append("\n")
            headers = [cell.text.strip().replace("\n", " ") for cell in table.rows[0].cells]
            md_lines.append("| " + " | ".join(headers) + " |")
            md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")

            for row in table.rows[1:]:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                md_lines.append("| " + " | ".join(row_cells) + " |")
            md_lines.append("\n")

        markdown_text = "\n\n".join(md_lines).strip()
        if not markdown_text:
            raise ValueError("DOCX document appears to have no readable text content.")

        logger.info(f"Extracted {len(markdown_text)} characters from DOCX document")
        return markdown_text

    except Exception as e:
        logger.error(f"Failed to convert DOCX to markdown: {e}")
        raise ValueError(f"Failed to parse DOCX document: {str(e)}") from e
